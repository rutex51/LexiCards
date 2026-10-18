#!/usr/bin/env python3
"""Generuje pliki testowe dla importu (tests/fixtures).

Pliki powstają w niezależnych programach – openpyxl i LibreOffice – żeby testy odczytu
nie sprawdzały wyłącznie plików zapisanych przez samą aplikację. Skoroszyt z obrazami
„w komórce” (Excel 365) jest składany ręcznie według budowy plików zapisywanych przez Excel.

Uruchomienie: python3 tools/gen_fixtures.py          – wszystkie pliki
              python3 tools/gen_fixtures.py notes    – tylko pliki z notatkami (wersja 1.1)
Wymaga: openpyxl, Pillow oraz LibreOffice (polecenie soffice).
"""

import csv
import io
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Font
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "tests" / "fixtures"
SAMPLE_CSV = ROOT / "docs" / "samples" / "lexicards_sample.csv"

POLISH_ROWS = [
    ["Dziedzina", "Moduł", "Lekcja", "Pojęcie PL", "Pojęcie EN", "Opis PL", "Opis EN", "Przykład PL", "Przykład EN", "Tagi", "Trudność", "Notatki"],
    ["Prawo", "Spółki", "Organy", "zarząd", "management board", "Organ prowadzący sprawy spółki; reprezentuje ją na zewnątrz.",
     "The body that runs the company's affairs.", "Zarząd zatwierdził „plan naprawczy”.", 'The board approved the "recovery plan".', "prawo, spółki", "3", "źródło: KSH"],
    ["Prawo", "Spółki", "Organy", "rada nadzorcza", "supervisory board", "Organ nadzoru.\nPowołuje zarząd.", "", "", "", "", "", ""],
    ["Prawo", "Spółki", "Kapitał", "wkład niepieniężny", "contribution in kind", "", "", "Żółć, gęślą jaźń – test znaków.", "", "", "9", ""],
    ["Prawo", "Umowy", "Zawarcie", "oferta", "offer", "", "", "", "", "", "", ""],
]


def write_csv(path, rows, delimiter, encoding, bom=False):
    with open(path, "w", newline="", encoding=encoding) as handle:
        if bom:
            handle.write("﻿")
        csv.writer(handle, delimiter=delimiter, lineterminator="\r\n").writerows(rows)


def csv_fixtures():
    write_csv(OUT / "polish-windows-1250.csv", POLISH_ROWS, ";", "cp1250")
    # ISO-8859-2 nie ma cudzysłowów drukarskich ani półpauzy.
    plain = [[cell.replace("„", "'").replace("”", "'").replace("–", "-") for cell in row] for row in POLISH_ROWS]
    write_csv(OUT / "polish-iso-8859-2.csv", plain, ";", "iso8859_2")
    write_csv(OUT / "comma-utf8.csv", POLISH_ROWS, ",", "utf-8")
    write_csv(OUT / "tab-utf16.txt", POLISH_ROWS, "\t", "utf-16-le", bom=True)
    # Excel dopisuje czasem wiersz z separatorem.
    with open(OUT / "sep-directive.csv", "w", newline="", encoding="utf-8") as handle:
        handle.write("sep=|\r\n")
        csv.writer(handle, delimiter="|", lineterminator="\r\n").writerows(POLISH_ROWS[:3])
    # Bez nagłówków: dwie kolumny.
    write_csv(OUT / "no-header.csv", [["bilans", "balance sheet"], ["aktywa", "assets"], ["pasywa", "liabilities"]], ";", "utf-8")


def workbook_fixture():
    """Skoroszyt z kilkoma arkuszami i przypadkami brzegowymi."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Słownik"
    sheet.append(["Domain", "Module", "Lesson", "Term PL", "Term EN", "Description PL", "Difficulty"])
    sheet.append(["Metrology", "CMM", "Probes", "sonda", "probe", "Element stykający się z detalem.", 2])
    # Scalone komórki struktury: nazwa tylko w pierwszym wierszu grupy.
    sheet.append([None, None, None, "trzpień", "stylus", None, 4.0])
    sheet.append([None, None, "Accuracy", "wzorcowanie", "calibration", None, None])
    sheet.append([])  # pusty wiersz
    sheet.append(["Metrology", "GD&T", "Form", "płaskość", "flatness", "A < B & C > D \"cudzysłów\"", 1])
    sheet.merge_cells("A2:A4")
    sheet.merge_cells("B2:B4")
    # Tekst z formatowaniem części znaków (zapisywany jako kilka fragmentów).
    sheet["E7"] = CellRichText("straight", TextBlock(InlineFont(b=True), "ness"))
    sheet["D7"] = "prostoliniowość"
    sheet["A7"], sheet["B7"], sheet["C7"] = "Metrology", "GD&T", "Form"
    # Formuła zwracająca tekst i liczba ułamkowa.
    sheet["D8"] = "okrągłość"
    sheet["E8"] = '=CONCATENATE("round","ness")'
    sheet["F8"] = 0.125
    sheet["A8"], sheet["B8"], sheet["C8"] = "Metrology", "GD&T", "Form"
    sheet["A1"].font = Font(bold=True)

    notes = book.create_sheet("Notatki")
    notes.append(["To nie są dane do importu"])

    hidden = book.create_sheet("Ukryty")
    hidden.append(["Polski", "Angielski"])
    hidden.append(["ukryte", "hidden"])
    hidden.sheet_state = "hidden"

    empty = book.create_sheet("Pusty")
    assert empty.max_row == 1

    book.save(OUT / "openpyxl-cases.xlsx")


def libreoffice(source, target_format, out_dir, infilter=None):
    command = ["soffice", "--headless", "--convert-to", target_format, "--outdir", str(out_dir)]
    if infilter:
        command.append(f"--infilter={infilter}")
    command.append(str(source))
    subprocess.run(command, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)


def libreoffice_fixtures():
    """Te same dane przykładowe zapisane przez LibreOffice jako XLSX, XLS i ODS."""
    with tempfile.TemporaryDirectory() as temp:
        temp = Path(temp)
        source = temp / "libreoffice-sample.csv"
        shutil.copy(SAMPLE_CSV, source)
        # 59 = średnik, 34 = cudzysłów, 76 = UTF-8
        csv_filter = "Text - txt - csv (StarCalc):59,34,76,1"
        libreoffice(source, "xlsx", temp, csv_filter)
        shutil.copy(temp / "libreoffice-sample.xlsx", OUT / "libreoffice-sample.xlsx")
        libreoffice(source, "xls", temp, csv_filter)
        shutil.copy(temp / "libreoffice-sample.xls", OUT / "old-format.xls")
        libreoffice(source, "ods", temp, csv_filter)
        shutil.copy(temp / "libreoffice-sample.ods", OUT / "open-document.ods")
        # Skoroszyt z przypadkami brzegowymi przepisany przez LibreOffice: formuły dostają
        # zapamiętane wyniki, a teksty trafiają do wspólnej tablicy w innym układzie.
        cases = temp / "libreoffice-cases.xlsx"
        nested = temp / "converted"
        nested.mkdir()
        shutil.copy(OUT / "openpyxl-cases.xlsx", cases)
        libreoffice(cases, "xlsx", nested)
        shutil.copy(nested / "libreoffice-cases.xlsx", OUT / "libreoffice-cases.xlsx")


def picture(width, height, color, kind="PNG"):
    """Prosty obraz testowy: kolorowe tło i przekątna (format PNG albo JPEG)."""
    image = Image.new("RGB", (width, height), color)
    draw = ImageDraw.Draw(image)
    draw.line((0, 0, width, height), fill="white", width=6)
    draw.rectangle((8, 8, width - 9, height - 9), outline="black", width=3)
    buffer = io.BytesIO()
    image.save(buffer, kind, quality=85)
    return buffer.getvalue()


def notes_workbook():
    """Arkusz pojęć i arkusz notatek z formatowaniem i obrazami – zapisany przez openpyxl."""
    book = Workbook()
    terms = book.active
    terms.title = "Pojęcia"
    terms.append(["Domain", "Module", "Lesson", "Term PL", "Term EN"])
    terms.append(["Business English", "Finance", "Financial Statements", "bilans", "balance sheet"])
    terms.append(["Business English", "Finance", "Financial Statements", "aktywa", "assets"])

    notes = book.create_sheet("Notatki")
    notes.append(["Dziedzina", "Moduł", "Lekcja", "Notatka", "Obraz"])
    notes.append(["Business English", "Finance", "Financial Statements"])
    notes["D2"] = CellRichText(
        TextBlock(InlineFont(b=True, sz=14), "Bilans"),
        TextBlock(InlineFont(sz=11), "\nZestawienie "),
        TextBlock(InlineFont(i=True, sz=11), "aktywów"),
        TextBlock(InlineFont(sz=11), " i "),
        TextBlock(InlineFont(u="single", sz=11), "pasywów"),
        TextBlock(InlineFont(sz=11), ".\n\n• Aktywa = Pasywa\n• "),
        TextBlock(InlineFont(color="FFFF0000", sz=11), "zobowiązania"),
        TextBlock(InlineFont(sz=11), " i kapitał"),
    )
    first = SheetImage(io.BytesIO(picture(240, 140, "#2f7fe6")))
    first.anchor = "E2"
    notes.add_image(first)
    # Drugi obraz tej samej lekcji w następnym wierszu (bez nazw – jak w scalonych komórkach).
    second = SheetImage(io.BytesIO(picture(180, 120, "#2fa84f", "JPEG")))
    second.anchor = "E3"
    notes.add_image(second)
    notes.append([None, None, None, None])
    notes["A3"] = None
    notes.append(["Business English", "Finance", "Investments", "Akcje i **obligacje** to *papiery wartościowe*.\n1. akcje\n2. obligacje"])
    notes.column_dimensions["D"].width = 60
    notes.row_dimensions[2].height = 110
    notes.row_dimensions[3].height = 95
    book.save(OUT / "openpyxl-notes.xlsx")


def excel_cell_images():
    """Obrazy „Umieść w komórce” (Excel 365): metadata.xml i richData, bez rysunków na arkuszu.
    Kolejność wpisów w plikach pośrednich celowo nie pokrywa się z kolejnością komórek."""
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pkg_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    rd_ns = "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata"
    header = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'

    def inline(ref, text):
        return f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{text}</t></is></c>'

    sheet = (header + f'<worksheet xmlns="{main_ns}" xmlns:r="{rel_ns}"><sheetData>'
             + '<row r="1">' + inline("A1", "Domain") + inline("B1", "Module") + inline("C1", "Lesson") + inline("D1", "Note") + inline("E1", "Image") + '</row>'
             + '<row r="2">' + inline("A2", "Metrology") + inline("B2", "CMM") + inline("C2", "Probes") + inline("D2", "Sonda dotykowa – budowa")
             + '<c r="E2" t="e" vm="1"><v>#VALUE!</v></c></row>'
             + '<row r="3"><c r="E3" t="e" vm="2"><v>#VALUE!</v></c></row>'
             + '</sheetData></worksheet>')
    metadata = (header + f'<metadata xmlns="{main_ns}" xmlns:xlrd="{rd_ns}">'
                + '<metadataTypes count="1"><metadataType name="XLRICHVALUE" minSupportedVersion="120000" copy="1" pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1" coerce="1"/></metadataTypes>'
                + '<futureMetadata name="XLRICHVALUE" count="2">'
                + '<bk><extLst><ext uri="{3e2802c4-a4d2-4d8b-9148-e3be6c30e623}"><xlrd:rvb i="1"/></ext></extLst></bk>'
                + '<bk><extLst><ext uri="{3e2802c4-a4d2-4d8b-9148-e3be6c30e623}"><xlrd:rvb i="0"/></ext></extLst></bk>'
                + '</futureMetadata>'
                + '<valueMetadata count="2"><bk><rc t="1" v="1"/></bk><bk><rc t="1" v="0"/></bk></valueMetadata>'
                + '</metadata>')
    # vm=1 → valueMetadata[0] → futureMetadata[1] → rv[0] → rel 1 → image2.png (niebieski)
    # vm=2 → valueMetadata[1] → futureMetadata[0] → rv[1] → rel 0 → image1.jpeg (zielony)
    rich_values = (header + f'<rvData xmlns="{rd_ns}" count="2"><rv s="0"><v>1</v><v>5</v></rv><rv s="0"><v>0</v><v>5</v></rv></rvData>')
    structures = (header + f'<rvStructures xmlns="{rd_ns}" count="1"><s t="_localImage"><k n="_rvRel:LocalImageIdentifier" t="i"/><k n="CalcOrigin" t="i"/></s></rvStructures>')
    value_rels = (header + f'<richValueRels xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel" xmlns:r="{rel_ns}"><rel r:id="rId1"/><rel r:id="rId2"/></richValueRels>')
    value_rels_rels = (header + f'<Relationships xmlns="{pkg_ns}">'
                       + f'<Relationship Id="rId1" Type="{rel_ns}/image" Target="../media/image1.jpeg"/>'
                       + f'<Relationship Id="rId2" Type="{rel_ns}/image" Target="../media/image2.png"/></Relationships>')
    workbook = (header + f'<workbook xmlns="{main_ns}" xmlns:r="{rel_ns}"><sheets><sheet name="Notes" sheetId="1" r:id="rId1"/></sheets></workbook>')
    workbook_rels = (header + f'<Relationships xmlns="{pkg_ns}">'
                     + f'<Relationship Id="rId1" Type="{rel_ns}/worksheet" Target="worksheets/sheet1.xml"/>'
                     + f'<Relationship Id="rId2" Type="{rel_ns}/sheetMetadata" Target="metadata.xml"/>'
                     + '<Relationship Id="rId3" Type="http://schemas.microsoft.com/office/2022/10/relationships/richValueRel" Target="richData/richValueRel.xml"/>'
                     + '<Relationship Id="rId4" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValue" Target="richData/rdrichvalue.xml"/>'
                     + '<Relationship Id="rId5" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValueStructure" Target="richData/rdrichvaluestructure.xml"/>'
                     + '</Relationships>')
    content_types = (header + '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                     + '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                     + '<Default Extension="xml" ContentType="application/xml"/>'
                     + '<Default Extension="png" ContentType="image/png"/><Default Extension="jpeg" ContentType="image/jpeg"/>'
                     + '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                     + '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                     + '<Override PartName="/xl/metadata.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/>'
                     + '<Override PartName="/xl/richData/rdrichvalue.xml" ContentType="application/vnd.ms-excel.rdrichvalue+xml"/>'
                     + '<Override PartName="/xl/richData/rdrichvaluestructure.xml" ContentType="application/vnd.ms-excel.rdrichvaluestructure+xml"/>'
                     + '<Override PartName="/xl/richData/richValueRel.xml" ContentType="application/vnd.ms-excel.richvaluerel+xml"/>'
                     + '</Types>')
    root_rels = (header + f'<Relationships xmlns="{pkg_ns}"><Relationship Id="rId1" Type="{rel_ns}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    with zipfile.ZipFile(OUT / "excel-cell-images.xlsx", "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        archive.writestr("xl/worksheets/sheet1.xml", sheet)
        archive.writestr("xl/metadata.xml", metadata)
        archive.writestr("xl/richData/rdrichvalue.xml", rich_values)
        archive.writestr("xl/richData/rdrichvaluestructure.xml", structures)
        archive.writestr("xl/richData/richValueRel.xml", value_rels)
        archive.writestr("xl/richData/_rels/richValueRel.xml.rels", value_rels_rels)
        archive.writestr("xl/media/image1.jpeg", picture(160, 100, "#2fa84f", "JPEG"))
        archive.writestr("xl/media/image2.png", picture(120, 90, "#2f7fe6"))


def notes_fixtures():
    notes_workbook()
    excel_cell_images()
    # Ten sam skoroszyt notatek przepisany przez LibreOffice: tekst sformatowany we wspólnej
    # tablicy, obrazy jako rysunki zakotwiczone przy komórce (twoCellAnchor).
    with tempfile.TemporaryDirectory() as temp:
        temp = Path(temp)
        source = temp / "libreoffice-notes.xlsx"
        shutil.copy(OUT / "openpyxl-notes.xlsx", source)
        nested = temp / "converted"
        nested.mkdir()
        libreoffice(source, "xlsx", nested)
        shutil.copy(nested / "libreoffice-notes.xlsx", OUT / "libreoffice-notes.xlsx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    if sys.argv[1:] != ["notes"]:
        csv_fixtures()
        workbook_fixture()
        libreoffice_fixtures()
    notes_fixtures()
    for path in sorted(OUT.iterdir()):
        print(f"{path.stat().st_size:>8}  tests/fixtures/{path.name}")


if __name__ == "__main__":
    main()
